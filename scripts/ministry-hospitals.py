"""Liste officielle des hôpitaux publics (ministère de la Santé, data.gov.ma, ODbL) → JSON.

Usage :
  python3 scripts/ministry-hospitals.py <fichier.xlsx> [sortie.json]
  python3 scripts/ministry-hospitals.py <fichier.xlsx> --brut        (toutes les lignes, contrôle du format)
  python3 scripts/ministry-hospitals.py --depuis-brut <brut.json> [sortie.json]

Les lignes sont recopiées sans modification (nom, catégorie, région, délégation, commune).
Aucune position, aucun téléphone n'est ajouté ici : le rapprochement avec OpenStreetMap se fait
dans le collecteur (scripts/collect-health.ts), et un hôpital non rapproché reste « sans position ».
"""
import json
import sys
from datetime import date

JEU = "https://data.gov.ma/data/fr/dataset/la-liste-des-hopitaux"
FICHIER = (
    "https://data.gov.ma/data/fr/dataset/2932e8a4-272c-4101-80ef-85519de47e7c/resource/"
    "7ac78092-9ef1-4389-a54f-06ac95efae8e/download/repartition-des-hopitaux-par-region-et-province-2022.xlsx"
)
# Date de mise à jour indiquée sur la page du jeu de données (le fichier ne la contient pas).
MIS_A_JOUR = "2023-11-14"
ENTETE = ["Région", "Delegation", "Commune", "Etablissement hospitalier", "Catégorie"]


def xlsx_rows(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in r]
            if any(cells):
                rows.append(cells)
        out[ws.title] = rows
    return out


def structure(sheets):
    for title, rows in sheets.items():
        head = next((i for i, r in enumerate(rows) if [c.strip() for c in r[:5]] == ENTETE), None)
        if head is None:
            continue
        data = rows[head + 1 :]
        legend = {r[6].strip(): r[7].strip() for r in data if len(r) > 7 and r[6].strip() and r[7].strip()}
        hopitaux = []
        for r in data:
            region, delegation, commune, nom, cat = (c.strip() for c in r[:5])
            if not nom or not cat:
                continue
            hopitaux.append({
                "nom": nom,
                "categorie": cat,
                "categorieLibelle": legend.get(cat),
                "region": region,
                "delegation": delegation,
                "commune": commune,
            })
        if len(hopitaux) < 100:
            raise SystemExit(f"seulement {len(hopitaux)} hôpitaux lus dans « {title} » : format inattendu, fichier non remplacé")
        return {
            "source": "Ministère de la Santé et de la Protection sociale — « Liste des hôpitaux 2022 » (Liste nominative des établissements hospitaliers par catégorie 2022)",
            "jeuDeDonnees": JEU,
            "fichier": FICHIER,
            "licence": "Open Data Commons Open Database License (ODbL)",
            "donneesMisesAJourLe": MIS_A_JOUR,
            "extraitLe": date.today().isoformat(),
            "perimetre": "Tout le Maroc (lignes recopiées sans modification)",
            "hopitaux": hopitaux,
        }
    raise SystemExit("en-tête attendu introuvable : format du fichier modifié, fichier non remplacé")


def main():
    args = sys.argv[1:]
    if args and args[0] == "--depuis-brut":
        with open(args[1], encoding="utf-8") as f:
            sheets = json.load(f)
        rest = args[2:]
    else:
        sheets = xlsx_rows(args[0])
        rest = args[1:]
        if "--brut" in rest:
            json.dump(sheets, sys.stdout, ensure_ascii=False, indent=1)
            return
    out = structure(sheets)
    text = json.dumps(out, ensure_ascii=False, indent=2) + "\n"
    if rest:
        with open(rest[0], "w", encoding="utf-8") as f:
            f.write(text)
        print(f"{len(out['hopitaux'])} hôpitaux écrits dans {rest[0]}")
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
